"""Export the source-based deployment/security runbook; no backend processes start."""
from pathlib import Path
import json
import os
import re
import zipfile

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from markdown_it import MarkdownIt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/STracer_Deployment_Workflow_and_Data_Protection.md'
DOCX = SOURCE.with_suffix('.docx')
PDF = SOURCE.with_suffix('.pdf')
HTML = ROOT / 'artifacts/STracer_Deployment_Workflow_and_Data_Protection.html'
md = MarkdownIt('commonmark', {'html': False}).enable('table')
source = SOURCE.read_text(encoding='utf-8')
d = Document()
section = d.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.68)
section.left_margin = section.right_margin = Inches(.72)
normal = d.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.07
for style in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    d.styles[style].font.color.rgb = RGBColor.from_string('193C59')
    d.styles[style].paragraph_format.keep_with_next = True
header = section.header.paragraphs[0]
header.text = 'STRACER  |  DEPLOYMENT AND DATA PROTECTION'
header.style = 'Caption'
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Revision 2.0  |  30 October 2026  |  Page ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)


def rich(paragraph, text, small=False):
    bold = False
    italic = False
    hyperlink = None
    for token in md.parseInline(text)[0].children or []:
        if token.type == 'strong_open':
            bold = True
        elif token.type == 'strong_close':
            bold = False
        elif token.type == 'em_open':
            italic = True
        elif token.type == 'em_close':
            italic = False
        elif token.type == 'link_open':
            hyperlink = token.attrGet('href')
        elif token.type == 'link_close':
            hyperlink = None
        elif token.type in ['text', 'code_inline', 'softbreak']:
            content = token.content if token.type != 'softbreak' else ' '
            if hyperlink and hyperlink.startswith('https://'):
                link = OxmlElement('w:hyperlink')
                link.set(qn('r:id'), paragraph.part.relate_to(hyperlink, RT.HYPERLINK, is_external=True))
                r = OxmlElement('w:r')
                prop = OxmlElement('w:rPr')
                color = OxmlElement('w:color')
                color.set(qn('w:val'), '0563C1')
                prop.append(color)
                r.append(prop)
                t = OxmlElement('w:t')
                t.text = content
                r.append(t)
                link.append(r)
                paragraph._p.append(link)
            else:
                run = paragraph.add_run(content)
                run.bold, run.italic = bold, italic
                if small:
                    run.font.size = Pt(9)
                if token.type == 'code_inline':
                    run.font.name = 'Consolas'
                    run.font.size = Pt(8.2 if small else 9)


lines = source.splitlines()
i = 0
code = False
while i < len(lines):
    line = lines[i]
    stripped = line.strip()
    if stripped.startswith('```'):
        code = not code
    elif code:
        p = d.add_paragraph()
        p.paragraph_format.line_spacing = 1
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(8)
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'F0F4F8')
        p._p.get_or_add_pPr().append(shade)
    elif not stripped:
        pass
    elif stripped.startswith('#'):
        match = re.match(r'^(#{1,4})\s+(.*)$', stripped)
        if match:
            level = len(match.group(1))-1
            d.add_heading(match.group(2), level)
    elif stripped.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            cells = [x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?', x.replace(' ', '')) for x in cells):
                rows.append(cells)
            i += 1
        table = d.add_table(rows=1, cols=len(rows[0]))
        table.style = 'Light Shading Accent 1'
        repeat = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for cell, text in zip(table.rows[0].cells, rows[0]):
            rich(cell.paragraphs[0], text, small=True)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for row in rows[1:]:
            cells = table.add_row().cells
            for cell, text in zip(cells, row):
                rich(cell.paragraphs[0], text, small=True)
                cell.paragraphs[0].paragraph_format.space_after = Pt(4)
        d.add_paragraph()
        continue
    elif stripped.startswith('- [ ] '):
        rich(d.add_paragraph(style='List Bullet'), '☐ '+stripped[6:])
    elif re.match(r'^\d+\. ', stripped):
        rich(d.add_paragraph(), stripped)
    else:
        rich(d.add_paragraph(), stripped)
    i += 1
d.save(DOCX)

css = '''
@page { size:A4; margin:18mm 17mm 19mm; }
body { font-family:"Segoe UI",Arial,sans-serif; color:#203040; font-size:10pt; line-height:1.4; }
h1 { font-size:25pt; color:#193c59; line-height:1.15; }
h2 { font-size:15pt; color:#193c59; margin-top:20pt; }
h3 { font-size:12pt; color:#193c59; margin-top:15pt; }
h4 { font-size:10.5pt; margin-top:13pt; }
h1,h2,h3,h4 { break-after:avoid; }
p { margin:6pt 0; orphans:3; widows:3; }
table { border-collapse:collapse; width:100%; font-size:8.6pt; margin:10pt 0; table-layout:auto; }
thead { display:table-header-group; }
th { background:#e8eff6; text-align:left; }
th,td { border:1px solid #c8d4df; padding:6pt; vertical-align:top; overflow-wrap:anywhere; }
tr { break-inside:avoid; }
pre { background:#f0f4f8; padding:9pt; white-space:pre-wrap; overflow-wrap:anywhere; font-size:8pt; }
code { font-family:Consolas,monospace; font-size:0.9em; overflow-wrap:anywhere; }
li { margin:4pt 0; }
a { color:#165a8b; text-decoration:none; }
'''
HTML.parent.mkdir(exist_ok=True)
HTML.write_text('<!doctype html><html lang="en"><head><meta charset="utf-8"><title>STracer Deployment and Data Protection</title><style>'+css+'</style></head><body>'+md.render(source)+'</body></html>', encoding='utf-8')

from playwright.sync_api import sync_playwright
with sync_playwright() as p:
    browser = p.chromium.launch(channel='msedge', headless=True)
    page = browser.new_page()
    page.goto(HTML.as_uri(), wait_until='load')
    page.pdf(path=str(PDF), format='A4', print_background=True, prefer_css_page_size=True,
        display_header_footer=True, header_template='<div></div>',
        footer_template='<div style="font-family:Arial;font-size:8px;width:100%;text-align:center;color:#687787">STracer | Revision 2.0 | 30 October 2026 | Page <span class="pageNumber"></span> of <span class="totalPages"></span></div>')
    diagram = ROOT/'docs/backend-deployment-flow.svg'
    page.goto(diagram.as_uri(), wait_until='load')
    page.locator('svg').screenshot(path=str(diagram.with_suffix('.png')))
    browser.close()

from pypdf import PdfReader
reader = PdfReader(PDF)
pdf_text = '\n'.join(page.extract_text() or '' for page in reader.pages)
pdf_search_text = re.sub(r'(?<=-)\s*\n\s*', '', pdf_text)
for required in ['BCrypt', 'PasswordDeriveBytes', 'SHA-256', 'AES-CBC', 'Ace Cloud', '007', 'not explicitly configured']:
    assert required in source and required in pdf_search_text, required
check = Document(DOCX)
assert len(check.tables) >= 5
with zipfile.ZipFile(DOCX) as z:
    assert z.testzip() is None
secret_file = Path(os.environ['APPDATA'])/'Microsoft/UserSecrets/SignalTracker.LocalDevelopment/secrets.json'
def secret_values(obj, prefix=''):
    for key, value in obj.items():
        name = prefix+key
        if isinstance(value, dict):
            yield from secret_values(value, name+':')
        elif isinstance(value, str) and len(value) >= 12 and name.lower().endswith(('password','secretkey','apikey','secret','uploadtoken')):
            yield value
if secret_file.exists():
    private = json.loads(secret_file.read_text(encoding='utf-8-sig'))
    assert all(value not in source and value not in pdf_text for value in secret_values(private)), 'Credential included in document'
print(DOCX)
print(PDF)
print(f'Validated {len(reader.pages)} PDF pages, {len(check.tables)} Word tables, source coverage and no configured credential values.')
